import sys
from pathlib import Path
import tempfile
import unittest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.normalization import normalize_text
from core.document_loader import (
    validate_file,
    load_txt_text,
    load_csv_text,
    load_docx_text,
    load_pptx_text,
    load_document_text,
    DocumentParsingError,
)


class TestParsers(unittest.TestCase):
    def test_normalize_text(self):
        raw = "  Hello   \t  World!  \r\n\r\n\r\nNew paragraph. \x00 "
        cleaned = normalize_text(raw)
        self.assertEqual(cleaned, "Hello World!\n\nNew paragraph.")

    def test_validate_file_not_found(self):
        with self.assertRaises(FileNotFoundError):
            validate_file(Path("non_existent_file.pdf"))

    def test_validate_file_unsupported_extension(self):
        with tempfile.NamedTemporaryFile(suffix=".exe", delete=False) as f:
            f.write(b"content")
            temp_path = Path(f.name)
        try:
            with self.assertRaises(ValueError) as ctx:
                validate_file(temp_path)
            self.assertIn("Unsupported file format", str(ctx.exception))
        finally:
            temp_path.unlink(missing_ok=True)

    def test_validate_file_empty(self):
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as f:
            temp_path = Path(f.name)
        try:
            with self.assertRaises(ValueError) as ctx:
                validate_file(temp_path)
            self.assertIn("empty", str(ctx.exception))
        finally:
            temp_path.unlink(missing_ok=True)

    def test_load_txt_text(self):
        with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", encoding="utf-8", delete=False) as f:
            f.write("Line 1\nLine 2 with info.")
            temp_path = Path(f.name)
        try:
            pages = load_txt_text(temp_path)
            self.assertEqual(len(pages), 1)
            self.assertIn("Line 1", pages[0]["text"])
            self.assertIn("Line 2 with info.", pages[0]["text"])
        finally:
            temp_path.unlink(missing_ok=True)

    def test_load_csv_text(self):
        with tempfile.NamedTemporaryFile(suffix=".csv", mode="w", encoding="utf-8", delete=False) as f:
            f.write("Name,Role,Score\nAlice,Developer,95\nBob,Designer,90")
            temp_path = Path(f.name)
        try:
            pages = load_csv_text(temp_path)
            self.assertEqual(len(pages), 1)
            self.assertIn("Header: Name | Role | Score", pages[0]["text"])
            self.assertIn("Row 1: Alice | Developer | 95", pages[0]["text"])
        finally:
            temp_path.unlink(missing_ok=True)

    def test_load_docx_text(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("This is a test DOCX paragraph.")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Header1"
        table.rows[0].cells[1].text = "Header2"

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            doc.save(f.name)
            temp_path = Path(f.name)
        try:
            pages = load_docx_text(temp_path)
            self.assertEqual(len(pages), 1)
            self.assertIn("This is a test DOCX paragraph.", pages[0]["text"])
            self.assertIn("Header1 | Header2", pages[0]["text"])
        finally:
            temp_path.unlink(missing_ok=True)

    def test_load_pptx_text(self):
        from pptx import Presentation
        prs = Presentation()
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        tx_box = slide.shapes.add_textbox(0, 0, 100, 100)
        tf = tx_box.text_frame
        tf.text = "Slide 1 Content"

        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            prs.save(f.name)
            temp_path = Path(f.name)
        try:
            pages = load_pptx_text(temp_path)
            self.assertEqual(len(pages), 1)
            self.assertIn("Slide 1 Content", pages[0]["text"])
        finally:
            temp_path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
